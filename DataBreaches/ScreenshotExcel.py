import os
import math
import pandas as pd
from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

def excel_to_grouped_single_slide(excel_filename="data.xlsx", output_pptx_name="excel_grouped_slide.pptx", sheet_name=0):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(script_dir, excel_filename)

    if not os.path.exists(excel_path):
        print(f"Error: Could not find '{excel_filename}' in {script_dir}")
        return

    # 1. Read Excel and generate styled HTML
    df = pd.read_excel(excel_path, sheet_name=sheet_name)
    html_content = df.to_html(index=False, classes='excel-table')
    
    styled_html = f"""
    <html>
    <head>
        <style>
            body {{ 
                font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; 
                margin: 0; 
                padding: 20px;
                background-color: #ffffff;
            }}
            table.excel-table {{
                border-collapse: collapse;
                width: max-content;
                font-size: 13px;
                color: #333333;
            }}
            table.excel-table th, table.excel-table td {{
                border: 1px solid #d0d7de;
                padding: 8px 12px;
                text-align: left;
                white-space: nowrap;
            }}
            table.excel-table th {{
                background-color: #f6f8fa;
                font-weight: 600;
                position: sticky;
                top: 0;
                z-index: 10;
            }}
            table.excel-table tr:nth-child(even) {{
                background-color: #fcfcfc;
            }}
        </style>
    </head>
    <body>
        {html_content}
    </body>
    </html>
    """

    temp_html = os.path.join(script_dir, "temp_excel_view.html")
    with open(temp_html, "w", encoding="utf-8") as f:
        f.write(styled_html)

    temp_chunk_prefix = os.path.join(script_dir, "temp_chunk_")
    chunk_files = []

    # 2. Render and capture chunks safely using standard viewport shots
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        viewport_width = 1600
        viewport_height = 2000  # Safe viewport chunk height
        
        page = browser.new_page(
            viewport={"width": viewport_width, "height": viewport_height},
            device_scale_factor=1.5
        )
        page.goto(f"file://{temp_html}")
        
        total_height = page.evaluate("document.body.scrollHeight")
        num_chunks = math.ceil(total_height / viewport_height)
        
        print(f"Total layout height: {total_height}px. Capturing {num_chunks} clean chunks...")

        for i in range(num_chunks):
            y_offset = i * viewport_height
            page.evaluate(f"window.scrollTo(0, {y_offset});")
            page.wait_for_timeout(200)
            
            chunk_path = f"{temp_chunk_prefix}{i}.png"
            # Take a standard viewport screenshot at this scroll location
            page.screenshot(path=chunk_path)
            chunk_files.append(chunk_path)

        browser.close()

    # 3. Calculate total required slide height based on image proportions
    target_slide_width_inches = 12.0
    calculated_total_height_inches = 1.0  # Margins

    chunk_dimensions = []
    for cf in chunk_files:
        with Image.open(cf) as img:
            w, h = img.size
            img_height_inches = target_slide_width_inches * (h / w)
            chunk_dimensions.append((cf, img_height_inches))
            calculated_total_height_inches += img_height_inches

    # 4. Initialize PowerPoint presentation with a custom tall slide
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(min(56.0, max(7.5, calculated_total_height_inches)))
    
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 5. Import, position with math, and collect shapes for grouping
    print("Importing, positioning, and grouping chunks on the slide...")
    shape_objects = []
    current_top = Inches(0.5)
    left_pos = Inches(0.66)

    for chunk_path, img_h in chunk_dimensions:
        pic = slide.shapes.add_picture(chunk_path, left_pos, current_top, width=Inches(target_slide_width_inches))
        shape_objects.append(pic)
        current_top += Inches(img_h)

    # Group all picture shapes together into a single master object
    if shape_objects:
        slide.shapes.add_group_shape(shape_objects)

    # Save presentation
    output_pptx_path = os.path.join(script_dir, output_pptx_name)
    prs.save(output_pptx_path)

    # Clean up temporary chunk files and html
    for cf in chunk_files:
        if os.path.exists(cf):
            os.remove(cf)
    if os.path.exists(temp_html):
        os.remove(temp_html)

    print(f"Success! Grouped single-slide presentation saved to: {output_pptx_path}")

if __name__ == "__main__":
    excel_to_grouped_single_slide("data.xlsx", "excel_grouped_slide.pptx")